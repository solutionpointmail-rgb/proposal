"""
renewal_analysis.py — The Benefits Group client-facing Renewal Analysis (.docx, + PDF)

Builds the polished, client-facing analysis in the house style approved from the
Tennessee Men's Health (MMM, LLC) sample, Oct 2026: navy header band with a green
"RECOMMENDED" badge, navy section bars, tinted KPI tiles, tinted comparison tables,
a 2x2 Next Steps grid and a Bottom Line box.

The routine writes the CONTENT as JSON; this script owns all of the LAYOUT, so every
client's document looks identical no matter which run produced it.

Usage:
    python3 renewal_analysis.py content.json "Out Name.docx" [--pdf]

Content JSON (all text fields accept **bold** inline markup):
{
  "title": "Medical Renewal Analysis",
  "client": "Tennessee Men's Health (MMM, LLC)",
  "effective_date": "November 1, 2026",
  "badge": {"label": "RECOMMENDED", "lines": ["UHC BUY-DOWN", "$5,000"]},
  "sections": [
    {"heading": "RECOMMENDATION", "blocks": [
        {"type": "lead", "text": "We recommend ..."},
        {"type": "paragraph", "text": "**The important story** is ..."},
        {"type": "kpis", "items": [{"value": "$26,095.65", "label": "Current monthly cost", "tone": "blue"}]},
        {"type": "table", "columns": ["Option", "Monthly"], "widths": [2, 1],
         "rows": [{"cells": ["Current UHC", "$26,095.65"], "tone": "white"}]},
        {"type": "bullets", "items": ["...", "..."]},
        {"type": "closing", "text": "We do not believe ..."}
    ]}
  ],
  "next_steps": [{"title": "Confirm UHC Alternate", "text": "Use the ..."}],
  "bottom_line": "The 65% exact-plan renewal is not ...",
  "footer": "The Benefits Group  |  Benefits that feel like care."
}

Tones (KPI tiles and table rows): blue, red, green, gold, white.
KPI value colour follows the tone (blue->navy, red->red, green->green, gold->navy).
"""

import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

# ---- Brand palette (sampled from the approved MMM sample) -------------------------
NAVY = "17324D"
GREEN = "82B440"
GREEN_TEXT = "5E8C2A"
RED_TEXT = "B9534F"
TEXT = "263A47"
MUTED = "687986"
RULE = "D5DCE3"
TINT = {"blue": "EAF1F6", "red": "FBECEC", "green": "EEF6E8", "gold": "FFF4D8", "white": "FFFFFF"}
KPI_VALUE = {"blue": NAVY, "red": RED_TEXT, "green": GREEN_TEXT, "gold": NAVY, "white": NAVY}
STEP_TINTS = ["green", "blue", "gold", "green"]

FONT = "Cambria"          # Word: Cambria. LibreOffice PDF render: metric-compatible Caladea.
PAGE_W_IN = 8.5
MARGIN_IN = 0.6
BODY_W_IN = PAGE_W_IN - 2 * MARGIN_IN


def rgb(hex6):
    return RGBColor.from_string(hex6)


# ---- low-level XML helpers --------------------------------------------------------
def shade(cell, hex6):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6)
    tcPr.append(shd)


def cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right),
                      ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def table_borders(table, color=None, size=4, inside=True):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    edges = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{edge}")
        if color and edge in edges:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def fixed_layout(table, widths_in):
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths_in) * 1440)))
    tblW.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_in):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in):
            cell.width = Inches(w)


def row_no_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    if header:
        trPr.append(OxmlElement("w:tblHeader"))


def keep_table_together(table):
    """Keep every row with the next one so a short table never breaks across a page."""
    rows = table.rows
    for i, row in enumerate(rows):
        if i == len(rows) - 1:
            continue
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True


def para_fmt(p, before=0, after=0, line=1.15, align=None, keep_next=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if align is not None:
        p.alignment = align


def add_runs(p, text, size=10.5, color=TEXT, bold=False, italic=False):
    """Add text with **bold** inline markup."""
    parts = re.split(r"(\*\*.+?\*\*)", text or "")
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        run = p.add_run(part[2:-2] if is_bold else part)
        run.bold = bold or is_bold
        run.italic = italic
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
        run.font.name = FONT
    return p


def cell_text(cell, text, size=10, color=TEXT, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
              before=0, after=0):
    p = cell.paragraphs[0]
    para_fmt(p, before=before, after=after, line=1.1, align=align)
    add_runs(p, text, size=size, color=color, bold=bold)
    return p


# ---- building blocks ---------------------------------------------------------------
def header_band(doc, data):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t)
    fixed_layout(t, [BODY_W_IN - 1.9, 1.9])
    left, right = t.rows[0].cells
    for c, fill in ((left, NAVY), (right, GREEN)):
        shade(c, fill)
        cell_margins(c, top=200, bottom=200, left=220, right=220)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    cell_text(left, data.get("wordmark", "THE BENEFITS GROUP"), size=9.5, color="FFFFFF", bold=True, after=6)
    p = left.add_paragraph()
    para_fmt(p, after=8, line=1.0)
    add_runs(p, data.get("title", "Medical Renewal Analysis"), size=22, color="FFFFFF", bold=True)
    p = left.add_paragraph()
    para_fmt(p, line=1.0)
    sub = data.get("client", "")
    if data.get("effective_date"):
        sub += "   |   " + data["effective_date"]
    add_runs(p, sub, size=9.5, color="DCE4EC")

    badge = data.get("badge") or {}
    cell_text(right, badge.get("label", "RECOMMENDED"), size=8.5, color=NAVY, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    for i, line in enumerate(badge.get("lines", [])):
        p = right.add_paragraph()
        para_fmt(p, line=1.05, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_runs(p, line, size=15 if i == 0 else 14, color="FFFFFF", bold=True)
    spacer(doc, 6)


def spacer(doc, pts):
    p = doc.add_paragraph()
    para_fmt(p, line=1.0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("")
    r.font.size = Pt(max(1, pts))


def section_bar(doc, heading, fill=NAVY):
    """A shaded paragraph (not a table) so 'keep with next' can stop it orphaning at a page end."""
    p = doc.add_paragraph()
    para_fmt(p, before=10, after=8, line=1.0, keep_next=True)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)
    bdr = OxmlElement("w:pBdr")
    for edge, space in (("top", 5), ("bottom", 5), ("left", 6), ("right", 6)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), fill)
        bdr.append(el)
    pPr.append(bdr)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "120")
    ind.set(qn("w:right"), "120")
    pPr.append(ind)
    add_runs(p, heading.upper(), size=10.5, color="FFFFFF", bold=True)


def lead(doc, text):
    p = doc.add_paragraph()
    para_fmt(p, after=6, line=1.2)
    add_runs(p, text, size=11.5, color=NAVY, bold=True)


def paragraph(doc, text, keep_next=False):
    p = doc.add_paragraph()
    para_fmt(p, after=8, line=1.25, keep_next=keep_next)
    add_runs(p, text, size=10.5)


def closing(doc, text):
    p = doc.add_paragraph()
    para_fmt(p, before=2, after=8, line=1.2)
    add_runs(p, text, size=10.5, color=NAVY, bold=True)


def bullets(doc, items):
    for i, item in enumerate(items):
        p = doc.add_paragraph()
        para_fmt(p, after=3, line=1.2, keep_next=i < len(items) - 1)
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        add_runs(p, "•  ", size=10.5, color=GREEN_TEXT, bold=True)
        add_runs(p, item, size=10.5)
    spacer(doc, 4)


def kpis(doc, items):
    n = max(1, len(items))
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t, color="FFFFFF", size=24)
    fixed_layout(t, [BODY_W_IN / n] * n)
    row_no_split(t.rows[0])
    for cell, item in zip(t.rows[0].cells, items):
        tone = item.get("tone", "blue")
        shade(cell, TINT.get(tone, TINT["blue"]))
        cell_margins(cell, top=170, bottom=170, left=110, right=110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_text(cell, item.get("value", ""), size=16, color=KPI_VALUE.get(tone, NAVY), bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
        p = cell.add_paragraph()
        para_fmt(p, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_runs(p, item.get("label", ""), size=8.5, color=MUTED)
    spacer(doc, 8)


def table(doc, block):
    cols = block["columns"]
    rows = block.get("rows", [])
    weights = block.get("widths") or ([1.6] + [1] * (len(cols) - 1))
    total = float(sum(weights))
    widths = [BODY_W_IN * w / total for w in weights]
    t = doc.add_table(rows=1 + len(rows), cols=len(cols))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t, color=RULE, size=4)
    fixed_layout(t, widths)
    hdr = t.rows[0]
    row_no_split(hdr, header=True)
    for i, (cell, name) in enumerate(zip(hdr.cells, cols)):
        shade(cell, NAVY)
        cell_margins(cell, top=70, bottom=70, left=100, right=100)
        cell_text(cell, name, size=8.5, color="FFFFFF", bold=True,
                  align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    for r, row in enumerate(rows, start=1):
        tr = t.rows[r]
        row_no_split(tr)
        tone = row.get("tone", "white")
        emphasis = row.get("bold", False)
        for i, (cell, val) in enumerate(zip(tr.cells, row.get("cells", []))):
            shade(cell, TINT.get(tone, "FFFFFF"))
            cell_margins(cell, top=60, bottom=60, left=100, right=100)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            last_text = i == len(cols) - 1 and block.get("last_column_left", False)
            cell_text(cell, str(val), size=9, color=NAVY if (i == 0 or emphasis) else TEXT,
                      bold=(i == 0 or emphasis),
                      align=WD_ALIGN_PARAGRAPH.LEFT if (i == 0 or last_text) else WD_ALIGN_PARAGRAPH.CENTER)
    if len(rows) <= 8:
        keep_table_together(t)
    spacer(doc, 8)
    if block.get("note"):
        p = doc.add_paragraph()
        para_fmt(p, after=8, line=1.15)
        add_runs(p, block["note"], size=8.5, color=MUTED, italic=True)


def next_steps(doc, steps):
    if not steps:
        return
    section_bar(doc, "NEXT STEPS", fill=GREEN)
    # pull the bar tight to the grid, like the sample
    doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
    pairs = [steps[i:i + 2] for i in range(0, len(steps), 2)]
    t = doc.add_table(rows=len(pairs), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t, color="FFFFFF", size=12)
    fixed_layout(t, [BODY_W_IN / 2] * 2)
    n = 0
    for r, pair in enumerate(pairs):
        row = t.rows[r]
        row_no_split(row)
        for c in range(2):
            cell = row.cells[c]
            if c < len(pair):
                step = pair[c]
                n += 1
                shade(cell, TINT[STEP_TINTS[(n - 1) % len(STEP_TINTS)]])
                cell_margins(cell, top=140, bottom=160, left=150, right=150)
                cell_text(cell, f"{n}. {step.get('title', '')}", size=10, color=NAVY, bold=True, after=5)
                p = cell.add_paragraph()
                para_fmt(p, line=1.2)
                add_runs(p, step.get("text", ""), size=9)
            else:
                shade(cell, "FFFFFF")
    keep_table_together(t)
    spacer(doc, 12)


def bottom_line(doc, text):
    if not text:
        return
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t)
    fixed_layout(t, [BODY_W_IN])
    cell = t.rows[0].cells[0]
    row_no_split(t.rows[0])
    shade(cell, TINT["green"])
    cell_margins(cell, top=200, bottom=220, left=300, right=300)
    cell_text(cell, "BOTTOM LINE", size=9, color=GREEN_TEXT, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    p = cell.add_paragraph()
    para_fmt(p, line=1.3, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_runs(p, text, size=11.5, color=NAVY, bold=True)


def footer(doc, text):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_runs(p, text, size=8.5, color=MUTED)


# ---- main --------------------------------------------------------------------------
BLOCKS = {
    "lead": lambda d, b: lead(d, b["text"]),
    "paragraph": lambda d, b: paragraph(d, b["text"]),
    "closing": lambda d, b: closing(d, b["text"]),
    "bullets": lambda d, b: bullets(d, b["items"]),
    "kpis": lambda d, b: kpis(d, b["items"]),
    "table": table,
}


def build(data, out_path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(PAGE_W_IN), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(MARGIN_IN))
    sec.top_margin, sec.bottom_margin = Inches(0.55), Inches(0.6)
    sec.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(TEXT)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)

    doc.core_properties.title = f"{data.get('client', '')} — {data.get('title', 'Renewal Analysis')}"
    doc.core_properties.author = "The Benefits Group"

    header_band(doc, data)
    for section in data.get("sections", []):
        section_bar(doc, section["heading"])
        blocks = section.get("blocks", [])
        for i, block in enumerate(blocks):
            kind = block.get("type", "paragraph")
            if kind not in BLOCKS:
                raise ValueError(f"Unknown block type: {kind}")
            # A paragraph that introduces a table/kpi row should not be stranded above it.
            nxt = blocks[i + 1]["type"] if i + 1 < len(blocks) else None
            if kind == "paragraph" and nxt in ("table", "kpis", "bullets"):
                paragraph(doc, block["text"], keep_next=True)
            else:
                BLOCKS[kind](doc, block)
    next_steps(doc, data.get("next_steps", []))
    bottom_line(doc, data.get("bottom_line"))
    footer(doc, data.get("footer", "The Benefits Group  |  Benefits that feel like care."))
    doc.save(out_path)
    return out_path


def to_pdf(docx_path):
    out_dir = str(Path(docx_path).resolve().parent)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, str(docx_path)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    return str(Path(docx_path).with_suffix(".pdf"))


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    content = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    out = build(content, sys.argv[2])
    print(out)
    if "--pdf" in sys.argv[3:]:
        print(to_pdf(out))
